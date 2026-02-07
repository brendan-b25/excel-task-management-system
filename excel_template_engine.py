"""
Advanced Excel Template Creator with Multiple Sheets, Formulas, and Conditional Formatting
Author: brendan-b25
"""

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, numbers
from openpyxl.formatting.rule import ColorScaleRule, CellIsRule, FormulaRule
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.utils import get_column_letter
from datetime import datetime, timedelta
import pandas as pd
import json

class ExcelTemplateCreator:
    def __init__(self, filename='template.xlsx'):
        self.filename = filename
        self.wb = Workbook()
        self.wb.remove(self.wb.active)  # Remove default sheet
        self.sheets = {}
        
    def create_sheet(self, name, headers, column_widths=None):
        """Create a new sheet with styled headers"""
        ws = self.wb.create_sheet(name)
        self.sheets[name] = ws
        
        # Add headers
        ws.append(headers)
        
        # Style headers
        header_fill = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=12)
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        for idx, cell in enumerate(ws[1], 1):
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = border
            
            # Set column width
            if column_widths and idx <= len(column_widths):
                ws.column_dimensions[get_column_letter(idx)].width = column_widths[idx-1]
            else:
                ws.column_dimensions[get_column_letter(idx)].width = 15
        
        # Freeze header row
        ws.freeze_panes = 'A2'
        
        return ws
    
    def add_data_with_formulas(self, sheet_name, data, formulas=None):
        """Add data with complex formulas"""
        ws = self.sheets[sheet_name]
        
        for row_data in data:
            ws.append(row_data)
        
        # Add formulas if provided
        if formulas:
            for formula in formulas:
                row = formula['row']
                col = formula['col']
                formula_text = formula['formula']
                ws[f'{get_column_letter(col)}{row}'] = formula_text
    
    def add_conditional_formatting(self, sheet_name, range_str, rules):
        """Add conditional formatting with multiple rules"""
        ws = self.sheets[sheet_name]
        
        for rule_type, rule_config in rules.items():
            if rule_type == 'color_scale':
                # Traffic light color scale
                ws.conditional_formatting.add(
                    range_str,
                    ColorScaleRule(
                        start_type='min', start_color='F8696B',  # Red
                        mid_type='percentile', mid_value=50, mid_color='FFEB84',  # Yellow
                        end_type='max', end_color='63BE7B'  # Green
                    )
                )
            
            elif rule_type == 'highlight_above':
                # Highlight values above threshold
                ws.conditional_formatting.add(
                    range_str,
                    CellIsRule(
                        operator='greaterThan',
                        formula=[str(rule_config['value'])],
                        fill=PatternFill(start_color=rule_config['color'], end_color=rule_config['color'], fill_type='solid'),
                        font=Font(bold=True)
                    )
                )
            
            elif rule_type == 'highlight_below':
                # Highlight values below threshold
                ws.conditional_formatting.add(
                    range_str,
                    CellIsRule(
                        operator='lessThan',
                        formula=[str(rule_config['value'])],
                        fill=PatternFill(start_color=rule_config['color'], end_color=rule_config['color'], fill_type='solid'),
                        font=Font(bold=True, color='FFFFFF')
                    )
                )
            
            elif rule_type == 'duplicate':
                # Highlight duplicates
                ws.conditional_formatting.add(
                    range_str,
                    FormulaRule(
                        formula=[f'COUNTIF({range_str},{range_str.split(":")[0]})>1'],
                        fill=PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
                    )
                )
    
    def add_summary_formulas(self, sheet_name, summary_row, data_range):
        """Add summary formulas (SUM, AVERAGE, MAX, MIN, etc.)"""
        ws = self.sheets[sheet_name]
        
        formulas = {
            'A': f'="TOTAL"',
            'B': f'=SUM({data_range})',
            'C': f'=AVERAGE({data_range})',
            'D': f'=MAX({data_range})',
            'E': f'=MIN({data_range})',
            'F': f'=COUNT({data_range})'
        }
        
        for col, formula in formulas.items():
            cell = ws[f'{col}{summary_row}']
            cell.value = formula
            cell.font = Font(bold=True, size=11)
            cell.fill = PatternFill(start_color='D5D8DC', end_color='D5D8DC', fill_type='solid')
    
    def add_chart(self, sheet_name, chart_type, data_range, title, position='H2'):
        """Add charts (Bar, Line, Pie)"""
        ws = self.sheets[sheet_name]
        
        if chart_type == 'bar':
            chart = BarChart()
        elif chart_type == 'line':
            chart = LineChart()
        elif chart_type == 'pie':
            chart = PieChart()
        else:
            return
        
        chart.title = title
        chart.style = 10
        
        # Parse data range
        data = Reference(ws, range_string=data_range)
        chart.add_data(data, titles_from_data=True)
        
        ws.add_chart(chart, position)
    
    def create_dashboard_sheet(self, name='Dashboard'):
        """Create a comprehensive dashboard sheet"""
        ws = self.wb.create_sheet(name, 0)  # Make it first sheet
        self.sheets[name] = ws
        
        # Title
        ws['A1'] = 'LIVE DASHBOARD'
        ws['A1'].font = Font(bold=True, size=20, color='2C3E50')
        ws.merge_cells('A1:H1')
        ws['A1'].alignment = Alignment(horizontal='center')
        
        # Last Updated
        ws['A2'] = 'Last Updated:'
        ws['B2'] = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        ws['B2'].font = Font(italic=True, color='7F8C8D')
        
        # KPI Cards styling
        kpi_fill = PatternFill(start_color='3498DB', end_color='3498DB', fill_type='solid')
        kpi_font = Font(bold=True, size=14, color='FFFFFF')
        
        # KPI placeholders
        kpis = [
            ('A4', 'Total Records'),
            ('C4', 'Active Items'),
            ('E4', 'Success Rate'),
            ('G4', 'Revenue')
        ]
        
        for cell, label in kpis:
            ws[cell] = label
            ws[cell].fill = kpi_fill
            ws[cell].font = kpi_font
            ws[cell].alignment = Alignment(horizontal='center', vertical='center')
            
            # Merge cells for KPI card
            col = cell[0]
            next_col = chr(ord(col) + 1)
            ws.merge_cells(f'{cell}:{next_col}4')
            ws.merge_cells(f'{cell[0]}5:{next_col}5')
        
        return ws
    
    def add_data_validation(self, sheet_name, cell_range, validation_type, values):
        """Add data validation (dropdown lists, etc.)"""
        from openpyxl.worksheet.datavalidation import DataValidation
        
        ws = self.sheets[sheet_name]
        
        if validation_type == 'list':
            dv = DataValidation(type="list", formula1=f'"{",".join(values)}"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(cell_range)
    
    def add_pivot_table_data(self, sheet_name, data_df):
        """Add data suitable for pivot tables"""
        ws = self.sheets[sheet_name]
        
        # Write DataFrame to sheet
        for r_idx, row in enumerate(data_df.itertuples(index=False), 2):
            for c_idx, value in enumerate(row, 1):
                ws.cell(row=r_idx, column=c_idx, value=value)
    
    def protect_sheet(self, sheet_name, password=None):
        """Protect sheet with optional password"""
        ws = self.sheets[sheet_name]
        ws.protection.sheet = True
        if password:
            ws.protection.password = password
    
    def save(self):
        """Save the workbook"""
        self.wb.save(self.filename)
        print(f"✅ Excel template saved: {self.filename}")
        return self.filename

# Example usage
if __name__ == '__main__':
    creator = ExcelTemplateCreator('advanced_template.xlsx')
    
    # Create dashboard
    creator.create_dashboard_sheet()
    
    # Create data sheet
    ws = creator.create_sheet(
        'Sales Data',
        ['Date', 'Product', 'Quantity', 'Unit Price', 'Total', 'Status'],
        [15, 20, 12, 12, 15, 12]
    )
    
    # Add sample data
    data = [
        ['2024-01-01', 'Product A', 10, 100, '=C2*D2', 'Completed'],
        ['2024-01-02', 'Product B', 5, 200, '=C3*D3', 'Pending'],
        ['2024-01-03', 'Product A', 15, 100, '=C4*D4', 'Completed'],
    ]
    creator.add_data_with_formulas('Sales Data', data)
    
    # Add conditional formatting
    creator.add_conditional_formatting(
        'Sales Data',
        'E2:E100',
        {
            'color_scale': True,
            'highlight_above': {'value': 1000, 'color': '63BE7B'}
        }
    )
    
    creator.save()
    print("Template created successfully!")
